# 產生 AIPAR ETA 通用介紹簡報（widescreen 16:9）。
# 執行：python docs/presentations/build_pptx.py

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "AIPAR-ETA-intro.pptx"
FONT = "Microsoft JhengHei"

PAPER = RGBColor(0xF7, 0xF6, 0xF2)
INK = RGBColor(0x23, 0x18, 0x15)
MUTED = RGBColor(0x5D, 0x5A, 0x55)
GOLD = RGBColor(0xEA, 0xBF, 0x29)
BLUE = RGBColor(0x25, 0x9F, 0xC8)
DARK = RGBColor(0x13, 0x15, 0x18)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xDE, 0xD3)
ACCENT = RGBColor(0xB8, 0x86, 0x0B)

TOTAL = 15
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def set_font(run, size: int, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(f"{{{A_NS}}}{tag}")
        if el is None:
            rPr.append(parse_xml(f'<a:{tag} xmlns:a="{A_NS}" typeface="{FONT}"/>'))
        else:
            el.set("typeface", FONT)


def add_shape(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.adjustments[0] = 0.08
    return sh


def textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set("anchor", anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return tf


def bullets(slide, l, t, w, h, items, size=16, color=INK, gap=10):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = "•  " + item
        set_font(run, size, False, color)
    return tf


def notes(slide, body: str) -> None:
    slide.notes_slide.notes_text_frame.text = body


def picture(slide, name: str, l, t, max_w, max_h):
    path = ASSETS / name
    im = Image.open(path)
    pw, ph = im.size
    aspect = pw / ph
    w, h = max_w, max_w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))
    return w, h


def trim_png(name: str, bg: tuple[int, int, int], pad: int = 28) -> None:
    path = ASSETS / name
    im = Image.open(path).convert("RGB")
    px = im.load()
    w, h = im.size

    def used(x: int, y: int) -> bool:
        r, g, b = px[x, y]
        return abs(r - bg[0]) > 14 or abs(g - bg[1]) > 14 or abs(b - bg[2]) > 14

    top = next((y for y in range(h) if any(used(x, y) for x in range(0, w, 3))), 0)
    bot = next((y for y in range(h - 1, -1, -1) if any(used(x, y) for x in range(0, w, 3))), h - 1)
    left = next((x for x in range(w) if any(used(x, y) for y in range(0, h, 3))), 0)
    right = next((x for x in range(w - 1, -1, -1) if any(used(x, y) for y in range(0, h, 3))), w - 1)
    box = (
        max(0, left - pad),
        max(0, top - pad),
        min(w, right + pad),
        min(h, bot + pad),
    )
    im.crop(box).save(path)


def footer(slide, page: int) -> None:
    add_shape(slide, 0.55, 7.18, 12.2, 0.015, LINE)
    textbox(slide, 0.55, 7.22, 4.5, 0.24, "AIPAR ETA", 10, False, MUTED)
    textbox(slide, 5.0, 7.22, 4.0, 0.24, "實驗室伙食系統", 10, False, MUTED, PP_ALIGN.CENTER)
    textbox(slide, 10.4, 7.22, 2.3, 0.24, f"{page:02d}  /  {TOTAL:02d}", 10, False, MUTED, PP_ALIGN.RIGHT)


def content(prs, kicker: str, title: str, page: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, 0, 0, 13.333, 7.5, PAPER)
    add_shape(slide, 0, 0, 0.12, 7.5, GOLD)
    textbox(slide, 0.55, 0.28, 8.5, 0.28, kicker, 11, True, ACCENT)
    textbox(slide, 0.55, 0.52, 12.2, 0.55, title, 28, True, INK)
    footer(slide, page)
    return slide


def caption(slide, l, t, w, text: str) -> None:
    textbox(slide, l, t, w, 0.28, text, 10, False, MUTED)


def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def slide_01(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, 0, 0, 13.333, 7.5, DARK)
    add_shape(slide, 0, 0, 0.18, 7.5, GOLD)
    add_shape(slide, 0, 7.18, 13.333, 0.32, GOLD)
    picture(slide, "logo.png", 0.7, 1.15, 1.7, 1.5)
    textbox(slide, 0.7, 2.85, 11.5, 0.85, "AIPAR ETA", 54, True, WHITE)
    textbox(slide, 0.7, 3.7, 11.5, 0.45, "實驗室伙食系統", 24, False, GOLD)
    textbox(
        slide, 0.7, 4.4, 10.5, 0.8,
        "把餐廳菜單、論壇揪團與帳本收進同一套系統。\n寫入走 Discord，查閱走校內網站。",
        18, False, RGBColor(0xE9, 0xE7, 0xE2),
    )
    textbox(slide, 0.7, 6.55, 8, 0.35, "通用介紹  ·  2026-09-18  ·  台北時間", 12, False, RGBColor(0xA3, 0xA8, 0xB0))
    notes(slide, "開場：這是實驗室內部的伙食系統，不是對外點餐平台。等聽眾看完封面再進目錄。")


def slide_02(prs) -> None:
    slide = content(prs, "目錄", "今天會帶到這些", 2)
    items = [
        ("01", "這套系統在解決什麼"),
        ("02", "三件核心能力：菜單、揪團、帳本"),
        ("03", "一場午餐怎麼走完"),
        ("04", "Discord 與校內網站怎麼分工"),
        ("05", "技術棧與系統架構"),
        ("06", "金額、語言與使用方式"),
    ]
    for i, (num, label) in enumerate(items):
        col, row = i % 2, i // 2
        l, t = 0.55 + col * 6.2, 1.4 + row * 1.55
        add_round(slide, l, t, 5.8, 1.35, WHITE)
        textbox(slide, l + 0.25, t + 0.28, 1.1, 0.4, num, 20, True, GOLD)
        textbox(slide, l + 1.4, t + 0.42, 4.1, 0.5, label, 18, False, INK)
    notes(slide, "約 12–15 分鐘。強調聽完能回答：誰用、做什麼、錢怎麼算、從哪裡操作。")


def slide_03(prs) -> None:
    slide = content(prs, "背景", "實驗室中午，資訊不該散落在十則訊息裡", 3)
    cards = [
        ("要點什麼", "哪一家店、菜單是哪一版、有沒有人認得照片上的字。"),
        ("誰點了什麼", "論壇裡加點、改點、截止之後，清單要立刻是同一份。"),
        ("最後誰該付多少", "價格以庫內菜單為準，結算後知道要把錢拿給誰。"),
    ]
    for i, (title, body) in enumerate(cards):
        l = 0.55 + i * 4.15
        add_round(slide, l, 1.45, 3.95, 3.55, WHITE)
        add_shape(slide, l, 1.45, 3.95, 0.1, GOLD)
        textbox(slide, l + 0.28, 1.8, 3.4, 0.7, f"{i+1:02d}  {title}", 20, True, INK)
        textbox(slide, l + 0.28, 2.7, 3.4, 1.8, body, 15, False, MUTED)
    notes(slide, "用實驗室午餐的日常痛點帶入，不要談實作缺陷。聽眾若不是工程背景，停在這三句就夠。")


def slide_04(prs) -> None:
    slide = content(prs, "一句話", "AIPAR ETA 是實驗室內部的伙食系統", 4)
    add_round(slide, 0.55, 1.4, 12.2, 1.7, WHITE)
    textbox(
        slide, 0.85, 1.7, 11.6, 1.15,
        "Discord 建檔、揪團、點餐、結算；校內網站只負責查閱。\n金額由程式依庫內價格計算，模型不得自行定價。",
        20, False, INK,
    )
    facts = [
        ("給誰用", "實驗室成員。寫入在 Discord 伺服器，查閱在校內網路。"),
        ("兩個入口", "Bot 負責操作，網站負責看板。兩邊共用同一份 PostgreSQL。"),
        ("模型的角色", "幫忙讀菜單、聽懂自然語言。錢與帳一律由程式決定。"),
    ]
    for i, (title, body) in enumerate(facts):
        l = 0.55 + i * 4.15
        add_round(slide, l, 3.4, 3.95, 3.2, WHITE)
        textbox(slide, l + 0.25, 3.6, 3.45, 0.45, title, 16, True, ACCENT)
        textbox(slide, l + 0.25, 4.15, 3.45, 2.1, body, 15, False, INK)
    notes(slide, "這頁是整場的定位句，後面功能都回到這三行。可念一次 README 的副標題。")


def slide_05(prs) -> None:
    slide = content(prs, "功能總覽", "三件核心能力", 5)
    cards = [
        ("菜單建檔與查詢", "/餐廳 新增 建檔。\n/菜單 上傳 丟照片，或 /菜單 輸入 貼文字。\n辨識結果先當草稿，確認後才上線。"),
        ("揪團點餐", "/揪團 在論壇開貼文，自動貼菜單與彙總。\n按「點餐」選數量與品項，或直接打字。\n截止後顯示截止時刻，不再倒數。"),
        ("記帳與分攤", "/結算 把每人小計寫進帳本。\n/帳務 誰欠誰 給互抵後的欠款與最少轉帳建議。\n同場同人不會重複計費。"),
    ]
    for i, (title, body) in enumerate(cards):
        l = 0.55 + i * 4.15
        add_round(slide, l, 1.4, 3.95, 5.15, WHITE)
        add_shape(slide, l, 1.4, 0.12, 5.15, GOLD if i != 1 else BLUE)
        textbox(slide, l + 0.4, 1.7, 3.3, 0.4, f"0{i+1}", 14, True, GOLD if i != 1 else BLUE)
        textbox(slide, l + 0.4, 2.15, 3.3, 0.7, title, 20, True, INK)
        textbox(slide, l + 0.4, 2.95, 3.3, 3.2, body, 15, False, MUTED)
    notes(slide, "三塊對應產品範圍。下一頁起各用一張真實畫面展開。")


def slide_06(prs) -> None:
    slide = content(prs, "功能一", "菜單建檔與查詢", 6)
    bullets(slide, 0.55, 1.35, 5.7, 4.6, [
        "先登記餐廳，再為它建立菜單版本。",
        "照片辨識支援直書菜單；也可以人工貼文字。",
        "選用 OCR 對帳，預覽時標出需要人工核對的行。",
        "沒有人按下「確認寫入」，菜單不會上線。",
        "之後可用指令或 @ bot 問「這家有什麼」。",
    ], 15, INK, 12)
    add_round(slide, 6.35, 1.28, 6.45, 5.55, WHITE)
    picture(slide, "web-menu.png", 6.5, 1.4, 6.15, 5.05)
    caption(slide, 6.5, 6.55, 6.15, "校內網站實機截圖　·　雷荷豆腐鍋專賣店菜單 v1")
    notes(slide, "右圖是本機網站 2026-09-18 的真實菜單頁，來源標為 vision，代表圖片辨識後經人工確認上線。")


def slide_07(prs) -> None:
    slide = content(prs, "功能二", "揪團點餐", 7)
    bullets(slide, 0.55, 1.35, 5.7, 4.7, [
        "管理員先指定論壇頻道，開團可 ping 身分組。",
        "Bot 開論壇貼文，貼上菜單與持續更新的彙總。",
        "點餐面板可選 1–10 份，品項可複選。",
        "也可以打「我要一個豆腐鍋跟一份地瓜薯條」。",
        "要改就說「取消地瓜薯條」，或按「清除我的」。",
        "價格永遠取自已上線的菜單，不接受口頭報價。",
    ], 15, INK, 10)
    add_round(slide, 6.35, 1.28, 6.45, 5.55, DARK)
    picture(slide, "discord-summary.png", 6.45, 1.38, 6.25, 5.15)
    caption(slide, 6.45, 6.55, 6.2, "Discord 介面示意　·　字串、按鈕與標誌取自本專案")
    notes(slide, "Discord 客戶端無法在簡報機截到實際伺服器畫面，故用專案字串表與網站標誌拼出對等畫面；品項來自雷荷豆腐鍋真實菜單。")


def slide_08(prs) -> None:
    slide = content(prs, "功能三", "記帳與分攤", 8)
    add_round(slide, 0.55, 1.35, 12.2, 1.45, WHITE)
    textbox(slide, 0.85, 1.55, 11.6, 1.05,
            "結算時寫下每人應付，並記住這場要把錢拿給誰。\n餘額由帳本紀錄加總得出，不另存一個會飄的數字。",
            16, False, INK)
    cards = [
        ("/結算", "每人小計寫成帳本紀錄。同場同一個人只計一次。"),
        ("/帳務 我的", "看自己的結餘，以及要拿給誰、誰要還你。"),
        ("/帳務 誰欠誰", "全體互抵後的欠款，加上最少轉帳建議。"),
        ("/帳務 付款", "記下實際付款；指定「付給」誰才抵銷債務。"),
    ]
    for i, (title, body) in enumerate(cards):
        l = 0.55 + (i % 4) * 3.15
        add_round(slide, l, 3.1, 3.0, 3.4, WHITE)
        textbox(slide, l + 0.2, 3.3, 2.6, 0.7, title, 16, True, ACCENT)
        textbox(slide, l + 0.2, 4.1, 2.6, 2.0, body, 14, False, INK)
    notes(slide, "結餘與債務是兩件事：結餘是吃了多少付了多少；債務是該把錢拿給哪一位。不要在這頁談空帳本。")


def slide_09(prs) -> None:
    slide = content(prs, "使用路徑", "一場午餐，六個步驟", 9)
    picture(slide, "flow.png", 0.4, 1.25, 12.5, 5.5)
    notes(slide, "跟著 1 到 6 走一遍。強調第 3 步的人工確認與第 6 步的結算。網站頁面列在圖下方黑條。")


def slide_10(prs) -> None:
    slide = content(prs, "兩個入口", "Discord 負責寫入，網站負責查閱", 10)
    bullets(slide, 0.55, 1.35, 5.6, 4.8, [
        "沒有公有網域：網站走校內 IP 與埠。",
        "頁面是伺服器端渲染，樣式內嵌，不依賴外部字型或 CDN。",
        "儀表板、餐廳菜單、揪團紀錄、帳務、系統狀態。",
        "深淺色可切換；系統要求減少動態時，動畫全部關閉。",
        "寫入只留 Discord 一條授權路徑，網站維持唯讀。",
    ], 15, INK, 11)
    add_round(slide, 6.3, 1.28, 6.5, 5.55, WHITE)
    picture(slide, "web-dashboard.png", 6.45, 1.4, 6.2, 5.05)
    caption(slide, 6.45, 6.55, 6.2, "校內網站實機截圖　·　儀表板")
    notes(slide, "右圖為 2026-09-18 本機網站儀表板，已有兩間餐廳與兩場揪團。不要解讀數字為營運 KPI。")


def slide_11(prs) -> None:
    slide = content(prs, "技術棧", "實驗室機器跑得動的組合", 11)
    rows = [
        ("執行期", "Node.js 24，TypeScript 直接執行，沒有前端建置步驟"),
        ("資料", "PostgreSQL 18.6；金額以整數計算，顯示為新台幣元"),
        ("編排", "Docker Compose：postgres、web、bot 三個容器"),
        ("Discord", "discord.js 14；斜線指令、Embed、按鈕與下拉"),
        ("網站", "Node 內建 HTTP、伺服器端渲染；刻意不引網頁框架"),
        ("模型", "自寫免費 LLM 閘道，OpenAI 相容格式，不引 LLM SDK"),
        ("辨識", "視覺模型讀菜單圖；可選 PP-OCRv6 做版面對帳"),
        ("時區", "全程台北時間，日期 YYYY-MM-DD、時間 HH:MM:SS"),
    ]
    for i, (k, v) in enumerate(rows):
        col, row = i % 2, i // 2
        l, t = 0.55 + col * 6.25, 1.35 + row * 1.3
        add_round(slide, l, t, 6.05, 1.15, WHITE)
        textbox(slide, l + 0.25, t + 0.18, 1.6, 0.35, k, 13, True, ACCENT)
        textbox(slide, l + 1.85, t + 0.32, 3.95, 0.55, v, 14, False, INK)
    notes(slide, "技術棧講「為什麼這樣選」：校內機器、免費 API、無公有網域。不必展開套件版本爭議。")


def slide_12(prs) -> None:
    slide = content(prs, "架構", "兩個入口，一份資料", 12)
    picture(slide, "architecture.png", 0.35, 1.2, 12.6, 5.55)
    notes(slide, "指出 bot 與 web 都進 PostgreSQL，只有 bot 會叫 LLM。分層：shared → db → domain → llm → 入口。")


def slide_13(prs) -> None:
    slide = content(prs, "設計原則", "讓帳算得準、讓人看得懂", 13)
    cards = [
        ("價格只來自菜單", "模型負責指到哪一項、要幾份。應付金額由程式查菜單計算。"),
        ("草稿必須有人確認", "圖片辨識與文字解析都先產生草稿，「確認寫入」才上線。"),
        ("規則優先、模型墊底", "對得到的品名就不打模型；沒有金鑰時，按鈕與下拉仍可點餐。"),
        ("兩種語言一起提供", "Discord 語言為中文時用繁體中文，其餘用英式英文。"),
    ]
    for i, (title, body) in enumerate(cards):
        col, row = i % 2, i // 2
        l, t = 0.55 + col * 6.25, 1.4 + row * 2.55
        add_round(slide, l, t, 6.05, 2.35, WHITE)
        add_shape(slide, l, t, 0.12, 2.35, GOLD if i % 2 == 0 else BLUE)
        textbox(slide, l + 0.4, t + 0.35, 5.35, 0.5, title, 18, True, INK)
        textbox(slide, l + 0.4, t + 1.0, 5.35, 1.05, body, 15, False, MUTED)
    notes(slide, "這四條是產品承諾，不是檢討。若有人問「AI 會不會算錯錢」，答案是錢不經過模型。")


def slide_14(prs) -> None:
    slide = content(prs, "怎麼開始", "八支斜線指令", 14)
    cmds = [
        ("/餐廳", "建檔與查詢"),
        ("/菜單", "上傳、輸入、查看版本"),
        ("/揪團", "開論壇貼文"),
        ("/結算", "寫入帳本"),
        ("/帳務", "我的／總覽／誰欠誰／付款"),
        ("/說明", "五段使用說明"),
        ("/網站", "開啟校內網站"),
        ("/設定", "論壇頻道與通知身分組"),
    ]
    for i, (cmd, desc) in enumerate(cmds):
        col, row = i % 2, i // 2
        l, t = 0.55 + col * 3.05, 1.32 + row * 1.28
        add_round(slide, l, t, 2.9, 1.12, WHITE)
        textbox(slide, l + 0.18, t + 0.18, 2.55, 0.38, cmd, 16, True, ACCENT)
        textbox(slide, l + 0.18, t + 0.58, 2.55, 0.38, desc, 12, False, MUTED)
    add_round(slide, 6.7, 1.32, 6.05, 5.2, DARK)
    picture(slide, "discord-help.png", 6.8, 1.42, 5.85, 4.85)
    caption(slide, 6.8, 6.55, 5.9, "Discord 介面示意　·　/說明 與「開啟網站」")
    notes(slide, "左列是現行指令表。英文預設名：restaurant、menu、groupbuy、settle、ledger、help、website、setup。")


def slide_15(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, 0, 0, 13.333, 7.5, DARK)
    add_shape(slide, 0, 0, 0.18, 7.5, GOLD)
    picture(slide, "logo.png", 0.7, 1.5, 1.5, 1.3)
    textbox(slide, 0.7, 3.1, 12, 0.7, "同一份菜單，同一本帳。", 32, True, WHITE)
    textbox(
        slide, 0.7, 3.95, 11, 1.1,
        "菜單先確認再上線，點餐對得到品項，結算由程式算錢。\nDiscord 操作，校內網站查閱。",
        18, False, RGBColor(0xE9, 0xE7, 0xE2),
    )
    textbox(slide, 0.7, 5.5, 11, 0.4, "問題歡迎現在提。", 16, False, GOLD)
    textbox(slide, 0.7, 6.55, 11, 0.35, "AIPAR ETA  ·  實驗室伙食系統  ·  2026-09-18", 12, False, RGBColor(0xA3, 0xA8, 0xB0))
    notes(slide, "收束到三句：確認菜單、對到品項、程式算錢。留下時間問指令或網站怎麼開。")


def main() -> None:
    trim_png("architecture.png", (247, 246, 242), 36)
    trim_png("flow.png", (247, 246, 242), 36)
    prs = new_prs()
    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)
    slide_14(prs)
    slide_15(prs)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
